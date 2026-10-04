"""Build the nine-slide presentation for the final project."""

# pylint: disable=line-too-long,missing-function-docstring
# pylint: disable=too-many-arguments,too-many-positional-arguments,too-many-locals

from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

from analise import (
    BLUE,
    CORAL,
    CREAM,
    INK,
    MINT,
    TEAL,
    YELLOW,
    clean_data,
    read_raw_data,
    summary_metrics,
)


PROJECT_DIR = Path(__file__).resolve().parents[1]
ASSET_DIR = PROJECT_DIR / "assets"
CHART_DIR = PROJECT_DIR / "graficos"
OUTPUT = PROJECT_DIR / "docs" / "apresentacao" / "apresentacao_rota_em_dia.pptx"
FONT = "Noto Sans"


def rgb(hex_color: str) -> RGBColor:
    value = hex_color.lstrip("#")
    return RGBColor.from_string(value)


def set_background(slide, color: str) -> None:
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = rgb(color)


def add_text(
    slide,
    text: str,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 20,
    color: str = INK,
    bold: bool = False,
    align=PP_ALIGN.LEFT,
    valign=MSO_ANCHOR.TOP,
    margin: float = 0.04,
    font: str = FONT,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(margin)
    frame.margin_right = Inches(margin)
    frame.margin_top = Inches(margin)
    frame.margin_bottom = Inches(margin)
    frame.vertical_anchor = valign
    paragraph = frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.text = text
    paragraph.font.name = font
    paragraph.font.size = Pt(size)
    paragraph.font.bold = bold
    paragraph.font.color.rgb = rgb(color)
    return box


def add_multiline(
    slide,
    lines: list[str],
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    size: float = 18,
    color: str = INK,
    bullet: bool = False,
    spacing: float = 8,
):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.margin_left = Inches(0.08)
    frame.margin_right = Inches(0.08)
    for index, line in enumerate(lines):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = FONT
        paragraph.font.size = Pt(size)
        paragraph.font.color.rgb = rgb(color)
        paragraph.space_after = Pt(spacing)
        if bullet:
            paragraph.text = f"•  {line}"
    return box


def add_card(
    slide,
    x: float,
    y: float,
    w: float,
    h: float,
    *,
    color: str = CREAM,
    radius=True,
    line: str | None = None,
):
    shape_type = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
    shape = slide.shapes.add_shape(shape_type, Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(1.2)
    else:
        shape.line.fill.background()
    return shape


def add_picture_contain(slide, path: Path, x: float, y: float, w: float, h: float):
    with Image.open(path) as image:
        source_ratio = image.width / image.height
    box_ratio = w / h
    if source_ratio >= box_ratio:
        width = w
        height = w / source_ratio
        left = x
        top = y + (h - height) / 2
    else:
        height = h
        width = h * source_ratio
        left = x + (w - width) / 2
        top = y
    return slide.shapes.add_picture(
        str(path), Inches(left), Inches(top), width=Inches(width), height=Inches(height)
    )


def set_transparency(shape, percent: int = 80) -> None:
    """Set shape fill transparency through the DrawingML alpha channel."""
    solid_fill = shape.fill._xPr.solidFill  # pylint: disable=protected-access
    color_node = solid_fill.getchildren()[0]
    for alpha in color_node.findall("{http://schemas.openxmlformats.org/drawingml/2006/main}alpha"):
        color_node.remove(alpha)
    alpha = OxmlElement("a:alpha")
    alpha.set("val", str((100 - percent) * 1000))
    color_node.append(alpha)


def add_transparent_shape(slide, shape_type, x, y, w, h, color, *, rotation=0):
    """Add a decorative shape with 80 percent transparency."""
    shape = slide.shapes.add_shape(
        shape_type, Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(color)
    set_transparency(shape)
    shape.line.fill.background()
    shape.rotation = rotation
    return shape


def add_doodles(slide, *, light: bool = True) -> None:
    pale = CREAM if light else TEAL
    colors = [YELLOW, CORAL, BLUE, pale]
    doodles = [
        (MSO_SHAPE.OVAL, -0.22, -0.20, 0.74, 0.74, colors[0]),
        (MSO_SHAPE.ARC, 12.40, -0.18, 0.95, 0.95, colors[2]),
        (MSO_SHAPE.STAR_8_POINT, 12.48, 6.62, 0.70, 0.70, colors[1]),
        (MSO_SHAPE.MOON, -0.24, 6.55, 0.72, 0.72, colors[3]),
    ]
    for shape_type, x, y, w, h, color in doodles:
        add_transparent_shape(slide, shape_type, x, y, w, h, color)

    add_transparent_shape(slide, MSO_SHAPE.OVAL, 10.85, 0.92, 1.25, 1.25, YELLOW)
    add_transparent_shape(slide, MSO_SHAPE.RECTANGLE, 11.43, 1.13, 0.09, 0.48, INK)
    add_transparent_shape(slide, MSO_SHAPE.RECTANGLE, 11.45, 1.50, 0.36, 0.09, INK, rotation=20)
    add_transparent_shape(slide, MSO_SHAPE.TEAR, 2.05, 5.85, 0.42, 0.58, BLUE, rotation=20)
    add_transparent_shape(slide, MSO_SHAPE.TEAR, 2.55, 6.23, 0.33, 0.46, BLUE, rotation=20)
    add_transparent_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 10.55, 5.72, 1.82, 0.94, TEAL)
    add_transparent_shape(slide, MSO_SHAPE.RECTANGLE, 10.84, 5.90, 1.22, 0.33, BLUE)
    add_transparent_shape(slide, MSO_SHAPE.OVAL, 10.83, 6.47, 0.31, 0.31, INK)
    add_transparent_shape(slide, MSO_SHAPE.OVAL, 11.78, 6.47, 0.31, 0.31, INK)


def add_header(slide, title: str, number: str, *, color: str = INK) -> None:
    add_text(slide, number, 0.45, 0.35, 0.55, 0.35, size=13, color=color, bold=True)
    add_text(slide, title, 0.88, 0.26, 11.6, 0.7, size=27, color=color, bold=True)


def add_metric_card(slide, value: str, label: str, x: float, y: float, w: float, color: str):
    add_card(slide, x, y, w, 1.25, color=CREAM)
    add_text(slide, value, x + 0.14, y + 0.13, w - 0.28, 0.52, size=25, color=color, bold=True)
    add_text(slide, label, x + 0.14, y + 0.69, w - 0.28, 0.38, size=11.5, color=INK)


def add_slide_number(slide, current: int, total: int) -> None:
    """Add a consistent page number to the bottom-right corner."""
    pill = add_card(slide, 12.18, 7.03, 0.78, 0.28, color=CREAM, line=TEAL)
    pill.fill.fore_color.rgb = rgb(CREAM)
    add_text(
        slide,
        f"{current}/{total}",
        12.22,
        7.055,
        0.70,
        0.20,
        size=9,
        color=TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
        valign=MSO_ANCHOR.MIDDLE,
        margin=0,
    )


def add_bus_drawing(slide, x: float, y: float, scale: float = 1.0) -> None:
    body = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x), Inches(y), Inches(2.0 * scale), Inches(1.2 * scale),
    )
    body.fill.solid()
    body.fill.fore_color.rgb = rgb(TEAL)
    body.line.fill.background()
    window = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE,
        Inches(x + 0.24 * scale), Inches(y + 0.19 * scale),
        Inches(1.52 * scale), Inches(0.48 * scale),
    )
    window.fill.solid()
    window.fill.fore_color.rgb = rgb(BLUE)
    window.line.fill.background()
    for wheel_x in [x + 0.32 * scale, x + 1.48 * scale]:
        wheel = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(wheel_x), Inches(y + 0.98 * scale),
            Inches(0.34 * scale), Inches(0.34 * scale),
        )
        wheel.fill.solid()
        wheel.fill.fore_color.rgb = rgb(INK)
        wheel.line.fill.background()
    for lamp_x in [x + 0.24 * scale, x + 1.58 * scale]:
        lamp = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(lamp_x), Inches(y + 0.77 * scale),
            Inches(0.16 * scale), Inches(0.16 * scale),
        )
        lamp.fill.solid()
        lamp.fill.fore_color.rgb = rgb(YELLOW)
        lamp.line.fill.background()


def slide_cover(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, CREAM)
    add_doodles(slide)
    add_picture_contain(slide, ASSET_DIR / "logo-rota-em-dia.png", 0.05, 0.42, 5.35, 5.25)
    add_text(slide, "Rota em Dia", 5.25, 1.28, 7.0, 0.9, size=42, color=TEAL, bold=True)
    add_text(
        slide,
        "Por que os ônibus fretados\nestão chegando atrasados?",
        5.28,
        2.14,
        6.8,
        1.25,
        size=25,
        color=INK,
        bold=False,
    )
    add_card(slide, 5.25, 3.64, 6.7, 1.34, color=MINT)
    add_text(slide, "Fundamentos de IA e Programação", 5.55, 3.92, 6.1, 0.35, size=16, color=TEAL, bold=True)
    add_text(slide, "Projeto Integrador · Outubro de 2026", 5.55, 4.34, 6.1, 0.3, size=13, color=INK)
    add_text(
        slide,
        "Juliana Ballin Lima  ·  Fernanda de Oliveira da Costa  ·  Pedro Henrique Oliveira Dias",
        0.76,
        6.58,
        11.8,
        0.34,
        size=11.5,
        color=TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_picture_contain(slide, ASSET_DIR / "logo-uea.png", 10.42, 0.20, 1.35, 0.58)
    add_picture_contain(slide, ASSET_DIR / "logo-sialabs.png", 11.85, 0.12, 0.94, 0.74)


def slide_summary(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, MINT)
    add_doodles(slide, light=False)
    add_header(slide, "Sumário", "01", color=TEAL)
    items = [
        ("1", "O Problema", "Atrasos no transporte", CORAL),
        ("2", "As Perguntas", "Três hipóteses", YELLOW),
        ("3", "O Dataset", "2.463 viagens", BLUE),
        ("4", "O Achado", "R03 e R05 desde 06/04", CORAL),
        ("5", "Os Entregáveis", "Colab e relatório", CREAM),
    ]
    points = [(1.00, 2.05), (3.43, 3.18), (5.88, 2.05), (8.33, 3.18), (10.77, 2.05)]
    centers = [(x + 0.34, y + 0.34) for x, y in points]
    for start, end in zip(centers, centers[1:]):
        connector = slide.shapes.add_connector(
            MSO_CONNECTOR.STRAIGHT,
            Inches(start[0]),
            Inches(start[1]),
            Inches(end[0]),
            Inches(end[1]),
        )
        connector.line.color.rgb = rgb(TEAL)
        connector.line.width = Pt(5)

    for index, ((number, title, subtitle, color), (x, y)) in enumerate(zip(items, points)):
        stop = slide.shapes.add_shape(
            MSO_SHAPE.OVAL, Inches(x), Inches(y), Inches(0.68), Inches(0.68)
        )
        stop.fill.solid()
        stop.fill.fore_color.rgb = rgb(color)
        stop.line.color.rgb = rgb(TEAL)
        stop.line.width = Pt(2)
        add_text(
            slide,
            number,
            x + 0.08,
            y + 0.11,
            0.52,
            0.42,
            size=20,
            color=INK,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
            margin=0,
        )
        label_y = 1.22 if index % 2 == 0 else 4.05
        add_card(slide, x - 0.52, label_y, 1.72, 1.04, color=CREAM, line=TEAL)
        add_text(slide, title, x - 0.42, label_y + 0.15, 1.52, 0.30, size=13.5, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, subtitle, x - 0.42, label_y + 0.52, 1.52, 0.31, size=9.8, color=INK, align=PP_ALIGN.CENTER)

    add_text(
        slide,
        "Uma Investigação Guiada Pelas Evidências",
        3.20,
        5.55,
        6.90,
        0.50,
        size=20,
        color=TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    add_bus_drawing(slide, 9.95, 5.38, 0.82)


def slide_problem(prs: Presentation, metrics: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, CORAL)
    add_doodles(slide, light=False)
    add_header(slide, "O Problema", "02", color=INK)
    add_card(slide, 0.72, 1.22, 7.35, 3.55, color=CREAM)
    add_text(
        slide,
        "Colaboradores chegam atrasados à linha porque os ônibus fretados não cumprem o horário planejado.",
        1.08,
        1.58,
        6.65,
        1.12,
        size=24,
        color=TEAL,
        bold=True,
    )
    add_text(
        slide,
        "O RH precisa saber se a piora é geral, onde ela se concentra e qual explicação os dados realmente sustentam.",
        1.08,
        3.03,
        6.5,
        0.88,
        size=16,
        color=INK,
    )
    add_metric_card(slide, "8", "Rotas", 8.48, 1.30, 1.75, TEAL)
    add_metric_card(slide, "3", "Turnos", 10.39, 1.30, 1.75, CORAL)
    add_metric_card(slide, "≤ 5 min", "Limite de pontualidade", 8.48, 2.82, 3.66, TEAL)
    add_metric_card(slide, "15 min", "Margem até o início do turno", 8.48, 4.34, 3.66, CORAL)
    add_bus_drawing(slide, 1.05, 5.34, 0.9)
    add_text(
        slide,
        f"Pontualidade geral tratada: {metrics['punctuality']:.1f}%",
        3.25,
        5.66,
        4.7,
        0.48,
        size=18,
        color=INK,
        bold=True,
    )
    add_text(slide, "Fev · Mar · Abr · Mai", 8.78, 6.35, 3.2, 0.34, size=13, color=INK, bold=True, align=PP_ALIGN.CENTER)


def slide_questions(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, CREAM)
    add_doodles(slide)
    add_header(slide, "Perguntas e Hipóteses", "03")
    add_card(slide, 0.78, 1.12, 11.75, 1.18, color=TEAL)
    add_text(
        slide,
        "A queda é geral por causa da chuva ou está concentrada em rotas, turnos e períodos específicos?",
        1.15,
        1.43,
        11.0,
        0.58,
        size=21,
        color=CREAM,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    cards = [
        ("H1", "Chuva", "A chuva forte explica a piora de pontualidade.", BLUE),
        ("H2", "Corredor", "R03 e R05 enfrentam uma restrição compartilhada.", CORAL),
        ("H3", "Ocorrências", "Pane, trânsito e pneus explicam a mudança.", YELLOW),
    ]
    for index, (label, title, text, color) in enumerate(cards):
        x = 0.82 + index * 4.16
        add_card(slide, x, 2.75, 3.68, 2.85, color=color)
        add_text(slide, label, x + 0.25, 2.99, 0.72, 0.54, size=25, color=TEAL, bold=True)
        add_text(slide, title, x + 1.02, 3.05, 2.25, 0.45, size=18, color=TEAL, bold=True)
        add_text(slide, text, x + 0.30, 3.81, 3.05, 0.90, size=15, color=INK, align=PP_ALIGN.CENTER)
        add_text(slide, "Testar · Comparar · Limitar", x + 0.42, 5.03, 2.80, 0.25, size=10, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Não buscamos uma resposta escondida. Construímos a explicação mais defensável.", 1.8, 6.35, 9.8, 0.38, size=16, color=TEAL, bold=True, align=PP_ALIGN.CENTER)


def slide_dataset(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, TEAL)
    add_doodles(slide, light=False)
    add_header(slide, "Dataset e Qualidade", "04", color=CREAM)
    add_metric_card(slide, "2.463", "Linhas recebidas", 0.76, 1.23, 2.24, CORAL)
    add_metric_card(slide, "11", "Variáveis", 3.18, 1.23, 2.24, TEAL)
    add_metric_card(slide, "102", "Dias de operação", 5.60, 1.23, 2.24, BLUE)
    add_metric_card(slide, "2.448", "Viagens após deduplicar", 8.02, 1.23, 3.15, CORAL)
    add_card(slide, 0.76, 2.88, 5.61, 3.55, color=CREAM)
    add_text(slide, "O Que a Base Descreve", 1.04, 3.12, 4.9, 0.4, size=19, color=TEAL, bold=True)
    add_multiline(
        slide,
        [
            "Data, turno, rota, zona e empresa",
            "Horários previsto e realizado",
            "Atraso, passageiros e chuva",
            "Ocorrência registrada pelo motorista",
        ],
        1.04,
        3.73,
        4.92,
        2.2,
        size=15,
        bullet=True,
        spacing=11,
    )
    add_card(slide, 6.68, 2.88, 5.80, 3.55, color=MINT)
    add_text(slide, "O Que Precisou de Cuidado", 6.99, 3.12, 5.1, 0.4, size=19, color=TEAL, bold=True)
    issues = [
        ("15", "Duplicatas exatas"),
        ("145", "Chuvas ausentes"),
        ("12", "Atrasos divergentes"),
        ("6 + 5", "Atrasos e lotações inválidos"),
    ]
    for index, (value, label) in enumerate(issues):
        y = 3.78 + index * 0.57
        add_text(slide, value, 7.01, y, 0.9, 0.32, size=16, color=CORAL, bold=True, align=PP_ALIGN.RIGHT)
        add_text(slide, label, 8.10, y, 3.8, 0.32, size=14, color=INK)
    add_text(slide, "Base bruta preservada · regras reproduzíveis", 7.15, 6.02, 4.75, 0.26, size=11.5, color=TEAL, bold=True, align=PP_ALIGN.CENTER)


def slide_result(prs: Presentation, metrics: dict) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, CREAM)
    add_doodles(slide)
    add_header(slide, "Resultado Principal", "05")
    add_picture_contain(slide, CHART_DIR / "01_serie_rotas.png", 0.55, 1.07, 8.85, 5.85)
    add_card(slide, 9.62, 1.33, 3.05, 1.38, color=CORAL)
    add_text(slide, "06/04", 9.84, 1.56, 2.58, 0.48, size=28, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Início da ruptura", 9.84, 2.09, 2.58, 0.28, size=12, color=INK, align=PP_ALIGN.CENTER)
    add_card(slide, 9.62, 3.02, 3.05, 1.38, color=MINT)
    add_text(slide, f"{metrics['target_punctuality']:.0f}%", 9.84, 3.24, 2.58, 0.48, size=28, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Pontualidade R03/R05 · Dia", 9.76, 3.79, 2.72, 0.28, size=11.5, color=TEAL, align=PP_ALIGN.CENTER)
    add_card(slide, 9.62, 4.72, 3.05, 1.38, color=YELLOW)
    add_text(slide, f"{metrics['target_median']:.0f} min", 9.84, 4.95, 2.58, 0.48, size=28, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "Mediana de atraso", 9.84, 5.49, 2.58, 0.28, size=12, color=TEAL, align=PP_ALIGN.CENTER)


def slide_evidence(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, YELLOW)
    add_doodles(slide, light=False)
    add_header(slide, "Evidências e Conclusão", "06", color=TEAL)
    add_picture_contain(slide, CHART_DIR / "02_heatmap_rota_turno.png", 0.55, 1.06, 6.58, 5.78)
    cards = [
        ("Chuva Agrava", "Correlação moderada, mas o efeito aparece em toda a operação.", BLUE),
        ("Obra Localiza", "“Obra na via” surge em 07/04 somente em R03 e R05.", CORAL),
        ("Noite Preserva", "As duas rotas mantêm 96% de pontualidade no turno noturno.", MINT),
    ]
    for index, (title, text, color) in enumerate(cards):
        y = 1.23 + index * 1.48
        add_card(slide, 7.27, y, 5.45, 1.18, color=color)
        add_text(slide, title, 7.55, y + 0.16, 1.65, 0.37, size=16, color=TEAL, bold=True)
        add_text(slide, text, 9.13, y + 0.15, 3.25, 0.70, size=12.5, color=INK)
    add_card(slide, 7.27, 5.72, 5.45, 0.88, color=TEAL)
    add_text(
        slide,
        "Hipótese mais plausível: restrição diurna no corredor compartilhado.",
        7.54,
        5.95,
        4.92,
        0.42,
        size=16,
        color=CREAM,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def slide_deliverables(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, BLUE)
    add_doodles(slide, light=False)
    add_header(slide, "Entregáveis Disponíveis", "07", color=INK)
    add_text(
        slide,
        "A análise pode ser reproduzida no Colab e consultada em um relatório técnico descritivo.",
        1.02,
        0.98,
        11.20,
        0.55,
        size=17,
        color=INK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )

    cards = [
        (
            0.83,
            "Notebook no Google Colab",
            "Código executável, diagnóstico, limpeza, gráficos e conclusão reproduzível.",
            "Abrir Notebook no Colab",
            "https://colab.research.google.com/github/JulianaBallin/TransitTrace/blob/main/transporte_fretado_rota_em_dia.ipynb",
            MINT,
        ),
        (
            6.86,
            "Relatório Técnico Descritivo",
            "Método, evidências, limitações, tabelas, recomendações e conclusão detalhada.",
            "Abrir Relatório Técnico",
            "https://github.com/JulianaBallin/TransitTrace/blob/main/docs/relatorio/relatorio_tecnico_rota_em_dia.pdf",
            YELLOW,
        ),
    ]
    for x, title, description, link_label, address, accent in cards:
        add_card(slide, x, 1.78, 5.63, 4.62, color=CREAM, line=TEAL)
        add_card(slide, x + 0.28, 2.08, 0.82, 0.82, color=accent)
        icon = "</>" if x < 2 else "PDF"
        add_text(
            slide,
            icon,
            x + 0.36,
            2.27,
            0.66,
            0.30,
            size=13,
            color=INK,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
            margin=0,
        )
        add_text(slide, title, x + 1.28, 2.18, 3.92, 0.63, size=20, color=TEAL, bold=True)
        add_text(slide, description, x + 0.42, 3.24, 4.78, 1.08, size=15, color=INK, align=PP_ALIGN.CENTER)
        link = add_card(slide, x + 0.82, 4.72, 4.00, 0.72, color=TEAL)
        link_text = add_text(
            slide,
            link_label,
            x + 1.00,
            4.91,
            3.64,
            0.28,
            size=13.5,
            color=CREAM,
            bold=True,
            align=PP_ALIGN.CENTER,
            valign=MSO_ANCHOR.MIDDLE,
            margin=0,
        )
        link_text.click_action.hyperlink.address = address
        link.click_action.hyperlink.address = address
        add_text(slide, "Disponível no Repositório TransitTrace", x + 0.78, 5.72, 4.10, 0.29, size=10.5, color=TEAL, bold=True, align=PP_ALIGN.CENTER)

    add_text(
        slide,
        "O Link do Relatório Também Está no Início do Notebook",
        2.34,
        6.62,
        8.64,
        0.32,
        size=14,
        color=INK,
        bold=True,
        align=PP_ALIGN.CENTER,
    )


def slide_notebook(prs: Presentation) -> None:
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_background(slide, CREAM)
    add_doodles(slide)
    add_header(slide, "Demonstração do Notebook", "08")
    add_card(slide, 0.70, 1.13, 7.52, 5.88, color="#F2F5F5", line=TEAL)
    add_card(slide, 0.94, 1.42, 7.04, 0.50, color=TEAL)
    add_text(slide, "Rota em Dia · Google Colab", 1.19, 1.55, 4.2, 0.22, size=12, color=CREAM, bold=True)
    steps = [
        ("1", "Carregar", "CSV + diagnóstico"),
        ("2", "Investigar", "Estatística + hipóteses"),
        ("3", "Limpar", "Pipeline + validação"),
        ("4", "Explicar", "5 gráficos + conclusão"),
    ]
    for index, (number, title, detail) in enumerate(steps):
        y = 2.20 + index * 1.03
        add_card(slide, 1.10, y, 6.68, 0.78, color=CREAM, line="#D0DADA")
        add_text(slide, number, 1.31, y + 0.15, 0.48, 0.38, size=19, color=CORAL, bold=True, align=PP_ALIGN.CENTER)
        add_text(slide, title, 1.98, y + 0.15, 1.72, 0.33, size=15, color=TEAL, bold=True)
        add_text(slide, detail, 3.65, y + 0.16, 3.62, 0.33, size=13, color=INK)
    add_card(slide, 8.55, 1.20, 4.12, 2.08, color=MINT)
    add_text(slide, "Executar Tudo", 8.91, 1.56, 3.42, 0.45, size=24, color=TEAL, bold=True, align=PP_ALIGN.CENTER)
    add_text(slide, "16 células de código\n20 células de texto", 8.91, 2.25, 3.42, 0.62, size=14, color=INK, align=PP_ALIGN.CENTER)
    add_card(slide, 8.55, 3.63, 4.12, 2.45, color=CORAL)
    add_text(slide, "Próximos Passos", 8.89, 3.92, 3.45, 0.40, size=20, color=INK, bold=True, align=PP_ALIGN.CENTER)
    add_multiline(
        slide,
        ["Cruzar cronograma de obras", "Coletar GPS por trecho", "Testar desvio ou antecipação"],
        8.93,
        4.53,
        3.30,
        1.28,
        size=13,
        color=INK,
        bullet=True,
        spacing=6,
    )
    add_text(slide, "Perguntas?", 9.15, 6.53, 3.0, 0.38, size=20, color=TEAL, bold=True, align=PP_ALIGN.CENTER)


def build_slides() -> Path:
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    raw = read_raw_data()
    data, _ = clean_data(raw)
    metrics = summary_metrics(data)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    prs.core_properties.title = "Rota em Dia · Transporte fretado"
    prs.core_properties.subject = "Projeto integrador de Fundamentos de IA e Programação"
    prs.core_properties.author = "Juliana Ballin Lima; Fernanda de Oliveira da Costa; Pedro Henrique Oliveira Dias"
    prs.core_properties.keywords = "transporte fretado, análise de dados, pontualidade"

    slide_cover(prs)
    slide_summary(prs)
    slide_problem(prs, metrics)
    slide_questions(prs)
    slide_dataset(prs)
    slide_result(prs, metrics)
    slide_evidence(prs)
    slide_deliverables(prs)
    slide_notebook(prs)

    total = len(prs.slides)
    for current, slide in enumerate(prs.slides, start=1):
        add_slide_number(slide, current, total)

    prs.save(OUTPUT)
    return OUTPUT


if __name__ == "__main__":
    print(build_slides())
